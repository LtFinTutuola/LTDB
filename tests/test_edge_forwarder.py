"""
tests/test_edge_forwarder.py
----------------------------
Unit and integration tests for the standalone LTDB Edge Forwarder
(scripts/excel_agent/ltdb_forwarder.py).
"""

import os
import json
import time
import openpyxl
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch, MagicMock
import urllib.error

import pytest

from scripts.excel_agent.ltdb_forwarder import (
    setup_logging,
    load_local_config,
    save_local_config,
    fetch_remote_config,
    cleanup_old_files,
    is_file_eligible_for_today,
    compute_row_hash,
    load_daily_state,
    save_daily_state,
    read_excel_rows,
    push_batch_to_ltdb,
    run_cycle
)


@pytest.fixture
def agent_dir(tmp_path):
    """Sets up a temporary agent directory with valid config.json."""
    config_file = tmp_path / "config.json"
    initial_config = {
        "ltdb_base_url": "http://mock-cloud:8000",
        "config_endpoint": "/api/v1/excel-synchronization/agent-config",
        "push_endpoint": "/api/v1/excel-synchronization/push-batch",
        "http_timeout_seconds": 5,
        "heartbeat_interval_seconds": 180,
        "fallback_config": {
            "excel_file_path": str(tmp_path / "mock_sales.xlsx"),
            "excel_code_column": "N",
            "start_row": 2,
            "excel_polling_interval": 30,
            "retention_days": 30
        }
    }
    with open(config_file, "w", encoding="utf-8") as f:
        json.dump(initial_config, f, indent=2)
    return tmp_path


def test_setup_logging(tmp_path):
    logger = setup_logging(tmp_path)
    assert logger is not None
    logs_dir = tmp_path / "logs"
    assert logs_dir.exists()
    today_str = date.today().isoformat()
    log_file = logs_dir / f"agent_{today_str}.log"
    logger.info("Test message")
    assert log_file.exists()


def test_load_and_save_local_config(agent_dir):
    cfg = load_local_config(agent_dir)
    assert cfg["ltdb_base_url"] == "http://mock-cloud:8000"

    cfg["ltdb_base_url"] = "http://new-cloud:9000"
    save_local_config(agent_dir, cfg)

    reloaded = load_local_config(agent_dir)
    assert reloaded["ltdb_base_url"] == "http://new-cloud:9000"


@patch("urllib.request.urlopen")
def test_fetch_remote_config_success(mock_urlopen, agent_dir):
    logger = setup_logging(agent_dir)
    local_cfg = load_local_config(agent_dir)

    mock_response = MagicMock()
    mock_response.status = 200
    mock_response.read.return_value = json.dumps({
        "status": "success",
        "data": {
            "excel_file_path": "C:\\Remote\\Sales.xlsx",
            "excel_code_column": "Z",
            "start_row": 1,
            "excel_polling_interval": 15,
            "retention_days": 60
        }
    }).encode("utf-8")
    mock_urlopen.return_value.__enter__.return_value = mock_response

    remote_cfg = fetch_remote_config(agent_dir, local_cfg, logger)
    assert remote_cfg["excel_code_column"] == "Z"
    assert remote_cfg["retention_days"] == 60

    # Verify fallback_config updated in local config.json
    reloaded = load_local_config(agent_dir)
    assert reloaded["fallback_config"]["excel_code_column"] == "Z"


@patch("urllib.request.urlopen")
def test_fetch_remote_config_offline_fallback(mock_urlopen, agent_dir):
    logger = setup_logging(agent_dir)
    local_cfg = load_local_config(agent_dir)

    # Simulate offline network error
    mock_urlopen.side_effect = urllib.error.URLError("Server unreachable")

    # Should fall back to local fallback_config (Alternative 1B)
    cfg = fetch_remote_config(agent_dir, local_cfg, logger)
    assert cfg["excel_code_column"] == "N"
    assert cfg["start_row"] == 2


def test_cleanup_old_files_retention(agent_dir):
    logger = setup_logging(agent_dir)
    state_dir = agent_dir / "state"
    state_dir.mkdir(parents=True, exist_ok=True)
    logs_dir = agent_dir / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)

    today = date.today()
    old_date = today - timedelta(days=40)
    recent_date = today - timedelta(days=10)

    # Create old and recent state files
    (state_dir / f"hashes_{old_date.isoformat()}.json").write_text("{}", encoding="utf-8")
    (state_dir / f"hashes_{recent_date.isoformat()}.json").write_text("{}", encoding="utf-8")

    # Create old and recent log files
    (logs_dir / f"agent_{old_date.isoformat()}.log").write_text("old log", encoding="utf-8")
    (logs_dir / f"agent_{recent_date.isoformat()}.log").write_text("recent log", encoding="utf-8")

    # Run cleanup with retention_days = 30
    cleanup_old_files(agent_dir, retention_days=30, logger=logger)

    # Old files should be deleted
    assert not (state_dir / f"hashes_{old_date.isoformat()}.json").exists()
    assert not (logs_dir / f"agent_{old_date.isoformat()}.log").exists()

    # Recent files should remain
    assert (state_dir / f"hashes_{recent_date.isoformat()}.json").exists()
    assert (logs_dir / f"agent_{recent_date.isoformat()}.log").exists()


def test_is_file_eligible_for_today(tmp_path):
    logger = setup_logging(tmp_path)
    excel_file = tmp_path / "test.xlsx"

    # Level 1: File does not exist
    assert is_file_eligible_for_today(excel_file, logger) is False

    # Create file modified now (today)
    excel_file.write_text("dummy", encoding="utf-8")
    assert is_file_eligible_for_today(excel_file, logger) is True

    # Level 2: Change mtime to yesterday
    yesterday_timestamp = (datetime.now() - timedelta(days=1)).timestamp()
    os.utime(str(excel_file), (yesterday_timestamp, yesterday_timestamp))
    assert is_file_eligible_for_today(excel_file, logger) is False


def test_compute_row_hash_determinism():
    row1 = {
        "excel_row_index": 5,
        "starting_price": 100.0,
        "selling_price": 90.0,
        "excel_file_column": 1,
        "is_exchange": False,
        "raw_article_code": "ART-123"
    }
    row2 = {
        "excel_row_index": 5,
        "starting_price": 100.0,
        "selling_price": 90.0,
        "excel_file_column": 1,
        "is_exchange": False,
        "raw_article_code": "  ART-123  "  # Spaces should be stripped
    }
    row3 = {
        "excel_row_index": 5,
        "starting_price": 100.0,
        "selling_price": 85.0,  # Changed price
        "excel_file_column": 1,
        "is_exchange": False,
        "raw_article_code": "ART-123"
    }

    hash1 = compute_row_hash(row1)
    hash2 = compute_row_hash(row2)
    hash3 = compute_row_hash(row3)

    assert hash1 == hash2
    assert hash1 != hash3


def test_daily_state_persistence(tmp_path):
    today = date(2026, 10, 20)
    state_dir = tmp_path / "state"

    # Initial load is empty
    state = load_daily_state(state_dir, today)
    assert state == {}

    # Save rows
    save_daily_state(state_dir, today, {"2": "hash2", "3": "hash3"})
    loaded = load_daily_state(state_dir, today)
    assert loaded["2"] == "hash2"
    assert loaded["3"] == "hash3"


def test_read_excel_rows(tmp_path):
    file_path = tmp_path / "test_read.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active

    # Row 1: Header / row 1
    sheet.cell(row=1, column=1, value="PREZZO")
    # Row 2: Normal sale (Col A=120, Col B=100 (Cat 0), Col N=CODE1)
    sheet.cell(row=2, column=1, value=120.0)
    sheet.cell(row=2, column=2, value=100.0)
    sheet.cell(row=2, column=14, value="CODE1")
    # Row 3: Exchange return (Col A=50, Col C=45 (Cat 1), Col L=X)
    sheet.cell(row=3, column=1, value=50.0)
    sheet.cell(row=3, column=3, value=45.0)
    sheet.cell(row=3, column=12, value="X")
    # Row 4: Empty row -> should stop stream
    wb.save(str(file_path))
    wb.close()

    rows = read_excel_rows(file_path, start_row=2, code_column="N")
    assert len(rows) == 2
    r2 = rows[0]
    assert r2["excel_row_index"] == 2
    assert r2["starting_price"] == 120.0
    assert r2["selling_price"] == 100.0
    assert r2["excel_file_column"] == 0
    assert r2["raw_article_code"] == "CODE1"
    assert r2["is_exchange"] is False

    r3 = rows[1]
    assert r3["excel_row_index"] == 3
    assert r3["starting_price"] == 50.0
    assert r3["selling_price"] == 45.0
    assert r3["excel_file_column"] == 1
    assert r3["is_exchange"] is True


@patch("urllib.request.urlopen")
def test_push_batch_to_ltdb_success_and_conflict(mock_urlopen, tmp_path):
    logger = setup_logging(tmp_path)
    today = date(2026, 10, 20)

    # 1. Success 200 OK
    mock_resp = MagicMock()
    mock_resp.status = 200
    mock_resp.read.return_value = json.dumps({
        "status": "success",
        "confirmed_row_indexes": [2, 3]
    }).encode("utf-8")
    mock_urlopen.return_value.__enter__.return_value = mock_resp

    confirmed = push_batch_to_ltdb(
        base_url="http://mock-cloud:8000",
        push_endpoint="/api/v1/excel-synchronization/push-batch",
        target_date=today,
        rows=[{"excel_row_index": 2}, {"excel_row_index": 3}],
        timeout=5,
        logger=logger
    )
    assert confirmed == [2, 3]

    # 2. 409 Conflict (historical pattern rejected)
    mock_urlopen.side_effect = urllib.error.HTTPError(
        url="http://mock-cloud:8000/push-batch",
        code=409,
        msg="Conflict - Historical file detected",
        hdrs={},
        fp=None
    )
    rejected = push_batch_to_ltdb(
        base_url="http://mock-cloud:8000",
        push_endpoint="/api/v1/excel-synchronization/push-batch",
        target_date=today,
        rows=[{"excel_row_index": 2}],
        timeout=5,
        logger=logger
    )
    assert rejected is None


@patch("scripts.excel_agent.ltdb_forwarder.push_batch_to_ltdb")
def test_run_cycle_delta_and_ack(mock_push, agent_dir):
    logger = setup_logging(agent_dir)
    local_cfg = load_local_config(agent_dir)
    runtime_cfg = local_cfg["fallback_config"]

    excel_file = Path(runtime_cfg["excel_file_path"])
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.cell(row=2, column=1, value=100.0)
    sheet.cell(row=2, column=2, value=90.0)
    wb.save(str(excel_file))
    wb.close()

    today = date.today()
    mock_push.return_value = [2]  # Confirmed row 2

    # Cycle 1: Row 2 is new -> push and record in state
    last_push = run_cycle(agent_dir, runtime_cfg, local_cfg, logger, last_push_time=0.0)
    assert mock_push.called
    assert len(mock_push.call_args[1]["rows"]) == 1

    state = load_daily_state(agent_dir / "state", today)
    assert "2" in state

    # Cycle 2: No modifications -> 0 rows pushed, no call to push_batch_to_ltdb
    mock_push.reset_mock()
    run_cycle(agent_dir, runtime_cfg, local_cfg, logger, last_push_time=time.time())
    assert not mock_push.called
