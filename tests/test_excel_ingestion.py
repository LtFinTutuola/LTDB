import pytest
from datetime import date
import openpyxl
import os
from tempfile import NamedTemporaryFile

from src.services.excel_ingestion_service import ingest_daily_sales
from src.repositories.excel_sales_repo import excel_sales_repo
from src.schemas.excel_sales import ExcelSaleStatus

def create_mock_excel(rows_data, has_headers_at_699=True):
    wb = openpyxl.Workbook()
    sheet = wb.active
    
    for r_idx, row_vals in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(row_vals, start=1):
            sheet.cell(row=r_idx, column=c_idx, value=val)
            
    if has_headers_at_699:
        # Put headers at 699
        headers = ["PREZZO", "CAT1", "CAT2", "CAT3", "CAT4", "CAT5", "CAT6", "CAT7", "CAT8", "J", "K", "RESO"]
        for c_idx, val in enumerate(headers, start=1):
            sheet.cell(row=699, column=c_idx, value=val)
            
    temp_file = NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    wb.close()
    return temp_file.name


def test_ingest_no_headers_format(db_session):
    # Rows: 1. normal, 2. normal exchange, 3. empty (should stop)
    # A=starting, B-I=selling, L=is_exchange
    rows = [
        [100.0, 90.0, None, None, None, None, None, None, None, None, None, None],
        [50.0, None, 45.0, None, None, None, None, None, None, None, None, "X"],
    ]
    file_path = create_mock_excel(rows)
    target_date = date(2026, 9, 29)
    
    try:
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 2
        
        # Verify first row
        records = excel_sales_repo.get_pending_records(db_session)
        assert len(records) == 2
        r1, r2 = records
        
        assert r1.excel_row_index == 1
        assert r1.excel_file_column == 'B'
        assert r1.starting_price == 100.0
        assert r1.selling_price == 90.0
        assert r1.is_exchange is False
        assert r1.status == ExcelSaleStatus.ORPHAN
        
        assert r2.excel_row_index == 2
        assert r2.excel_file_column == 'C'
        assert r2.starting_price == 50.0
        assert r2.selling_price == 45.0
        assert r2.is_exchange is True
        assert r2.status == ExcelSaleStatus.ORPHAN
    finally:
        os.remove(file_path)


def test_unprocessable_row(db_session):
    # Row with bad data
    rows = [
        ["invalid", 90.0, None, None, None, None, None, None, None, None, None, None],
    ]
    file_path = create_mock_excel(rows)
    target_date = date(2026, 9, 29)
    
    try:
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 1
        
        records = excel_sales_repo.get_pending_records(db_session)
        assert len(records) == 1
        r1 = records[0]
        
        assert r1.status == ExcelSaleStatus.UNPROCESSABLE
        assert r1.starting_price is None
        assert r1.selling_price is None
    finally:
        os.remove(file_path)


def test_row_count_resume(db_session):
    # Idempotency / resume
    rows = [
        [10.0, 9.0, None, None, None, None, None, None, None, None, None, None],
        [20.0, 19.0, None, None, None, None, None, None, None, None, None, None],
    ]
    file_path = create_mock_excel(rows)
    target_date = date(2026, 9, 29)
    
    try:
        # Ingest first time: should insert 2 rows
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 2
        
        # Second time: should insert 0 rows
        inserted_ids2 = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids2) == 0
        
        # Add a 3rd row and save
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        sheet.cell(row=3, column=1, value=30.0)
        sheet.cell(row=3, column=3, value=25.0) # Column C
        wb.save(file_path)
        wb.close()
        
        # Third time: should insert 1 row (the 3rd one)
        inserted_ids3 = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids3) == 1
        
        records = excel_sales_repo.get_pending_records(db_session)
        assert len(records) == 3
        
        r3 = next(r for r in records if r.excel_row_index == 3)
        assert r3.starting_price == 30.0
        assert r3.selling_price == 25.0
        assert r3.excel_file_column == 'C'
    finally:
        os.remove(file_path)
