import pytest
from datetime import date
import openpyxl
import os
from tempfile import NamedTemporaryFile

from src.services.excel_synchronization_service import ingest_daily_sales, find_candidates, reconcile_sale
from src.repositories.excel_synchronization_repo import excel_synchronization_repo
from src.schemas.excel_sales import ExcelSaleStatus

def create_mock_excel(rows_data, has_headers_at_699=True):
    wb = openpyxl.Workbook()
    sheet = wb.active
    
    for r_idx, row_vals in enumerate(rows_data, start=1):
        for c_idx, val in enumerate(row_vals, start=1):
            sheet.cell(row=r_idx, column=c_idx, value=val)
            
    if has_headers_at_699:
        # Put headers at 699
        headers = ["PREZZO", "CAT1", "CAT2", "CAT3", "CAT4", "CAT5", "CAT6", "CAT7", "CAT8", "J", "K", "RESO"]
        for c_idx, val in enumerate(headers, start=1):
            sheet.cell(row=699, column=c_idx, value=val)
            
    temp_file = NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(temp_file.name)
    wb.close()
    return temp_file.name


def test_ingest_no_headers_format(db_session):
    # Rows: 1. normal, 2. normal exchange, 3. empty (should stop)
    # A=starting, B-I=selling, L=is_exchange
    rows = [
        [100.0, 90.0, None, None, None, None, None, None, None, None, None, None],
        [50.0, None, 45.0, None, None, None, None, None, None, None, None, "X"],
    ]
    file_path = create_mock_excel(rows)
    target_date = date(2026, 9, 29)
    
    try:
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 2
        
        # Verify first row
        records = excel_synchronization_repo.get_pending_records(db_session)
        assert len(records) == 2
        r1, r2 = records
        
        assert r1.excel_row_index == 1
        assert r1.excel_file_column == 0
        assert r1.starting_price == 100.0
        assert r1.selling_price == 90.0
        assert r1.is_exchange is False
        assert r1.status == ExcelSaleStatus.ORPHAN
        
        assert r2.excel_row_index == 2
        assert r2.excel_file_column == 1
        assert r2.starting_price == 50.0
        assert r2.selling_price == 45.0
        assert r2.is_exchange is True
        assert r2.status == ExcelSaleStatus.ORPHAN
    finally:
        os.remove(file_path)


def test_unprocessable_row(db_session):
    # Row with bad data
    rows = [
        ["invalid", 90.0, None, None, None, None, None, None, None, None, None, None],
    ]
    file_path = create_mock_excel(rows)
    target_date = date(2026, 9, 29)
    
    try:
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 1
        
        records = excel_synchronization_repo.get_pending_records(db_session)
        assert len(records) == 1
        r1 = records[0]
        
        assert r1.status == ExcelSaleStatus.UNPROCESSABLE
        assert r1.starting_price is None
        assert r1.selling_price is None
    finally:
        os.remove(file_path)


def test_row_count_resume(db_session):
    # Idempotency / resume
    rows = [
        [10.0, 9.0, None, None, None, None, None, None, None, None, None, None],
        [20.0, 19.0, None, None, None, None, None, None, None, None, None, None],
    ]
    file_path = create_mock_excel(rows)
    target_date = date(2026, 9, 29)
    
    try:
        # Ingest first time: should insert 2 rows
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 2
        
        # Second time: should insert 0 rows
        inserted_ids2 = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids2) == 0
        
        # Add a 3rd row and save
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        sheet.cell(row=3, column=1, value=30.0)
        sheet.cell(row=3, column=3, value=25.0) # Column C
        wb.save(file_path)
        wb.close()
        
        # Third time: should insert 1 row (the 3rd one)
        inserted_ids3 = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids3) == 1
        
        records = excel_synchronization_repo.get_pending_records(db_session)
        assert len(records) == 3
        
        r3 = next(r for r in records if r.excel_row_index == 3)
        assert r3.starting_price == 30.0
        assert r3.selling_price == 25.0
        assert r3.excel_file_column == 1
    finally:
        os.remove(file_path)

from src.models.wms import Article, ArticleStatus, ArticleMovement, MovementReason, Batch, Supplier
from src.models.pim import ArticleBlueprint, Brand
from src.schemas.excel_sales import ExcelSaleCreate

def create_base_data(db_session):
    brand = Brand(name="Test Brand")
    db_session.add(brand)
    db_session.commit()
    
    bp = ArticleBlueprint(article_name="Test", description="Test desc", extended_description="Test ext desc", brand_id=str(brand.id))
    db_session.add(bp)
    db_session.commit()
    
    supp = Supplier(company_name="Test Supplier")
    db_session.add(supp)
    db_session.commit()
    
    batch = Batch(supplier_id=str(supp.id), delivery_note_number="123", document_date=date(2023,1,1))
    db_session.add(batch)
    db_session.commit()
    
    return str(bp.id), str(batch.id)

def test_find_candidates_and_reconcile(db_session):
    blueprint1_id, batch_id = create_base_data(db_session)
    
    a1 = Article(article_blueprint_id=blueprint1_id, batch_id=batch_id, supplier_code="VEND-XYZ", colors=["Red"], status=ArticleStatus.AVAILABLE)
    a2 = Article(article_blueprint_id=blueprint1_id, batch_id=batch_id, supplier_code="VEND-XYZ", colors=["Red"], status=ArticleStatus.AVAILABLE)
    a3 = Article(article_blueprint_id=blueprint1_id, batch_id=batch_id, supplier_code="VEND-XYZ", colors=["Blue"], status=ArticleStatus.AVAILABLE)
    a4 = Article(article_blueprint_id=blueprint1_id, batch_id=batch_id, supplier_code="VEND-XYZ", colors=["Blue"], status=ArticleStatus.AVAILABLE)
    db_session.add_all([a1, a2, a3, a4])
    db_session.commit()
    db_session.refresh(a1)
    
    candidates = find_candidates(db_session, "VEND-XYZ", is_exchange=False)
    assert len(candidates) == 2
    colors_returned = {c["colors"][0].lower() for c in candidates}
    assert "red" in colors_returned
    assert "blue" in colors_returned
    
    selected_article_id = str(a1.id)
    
    sale_in = ExcelSaleCreate(date=date(2023, 10, 1), excel_row_index=1, status=ExcelSaleStatus.ORPHAN, is_exchange=False)
    db_sale = excel_synchronization_repo.create(db_session, obj_in=sale_in)
    
    result = reconcile_sale(db_session, str(db_sale.id), selected_article_id)
    assert result["status"] == "ok"
    
    db_session.refresh(db_sale)
    assert db_sale.status == ExcelSaleStatus.RECONCILED
    assert str(db_sale.article_id) == selected_article_id
    
    db_session.refresh(a1)
    assert a1.status == ArticleStatus.SOLD
    
    mov = db_session.query(ArticleMovement).filter(ArticleMovement.article_id == a1.id).first()
    assert mov is not None
    assert mov.reason.code == "EXCEL_SALE"
    assert mov.reason.sign == -1

def test_reconcile_exchange(db_session):
    blueprint1_id, batch_id = create_base_data(db_session)
    
    a1 = Article(article_blueprint_id=blueprint1_id, batch_id=batch_id, supplier_code="VEND-RET", colors=["Green"], status=ArticleStatus.SOLD)
    db_session.add(a1)
    db_session.commit()
    db_session.refresh(a1)
    
    sale_in = ExcelSaleCreate(date=date(2023, 10, 2), excel_row_index=1, status=ExcelSaleStatus.ORPHAN, is_exchange=True)
    db_sale = excel_synchronization_repo.create(db_session, obj_in=sale_in)
    
    result = reconcile_sale(db_session, str(db_sale.id), str(a1.id))
    assert result["status"] == "ok"
    
    db_session.refresh(a1)
    assert a1.status == ArticleStatus.AVAILABLE
    
    mov = db_session.query(ArticleMovement).filter(ArticleMovement.article_id == a1.id).first()
    assert mov is not None
    assert mov.reason.code == "EXCEL_RETURN"
    assert mov.reason.sign == 1

def test_get_pending_sales(db_session):
    from src.services.excel_synchronization_service import get_pending_sales
    
    sale_in = ExcelSaleCreate(date=date(2023, 10, 5), excel_row_index=1, status=ExcelSaleStatus.ORPHAN, is_exchange=False)
    sale_in2 = ExcelSaleCreate(date=date(2023, 10, 5), excel_row_index=2, status=ExcelSaleStatus.UNPROCESSABLE, is_exchange=False)
    sale_in3 = ExcelSaleCreate(date=date(2023, 10, 6), excel_row_index=1, status=ExcelSaleStatus.ORPHAN, is_exchange=False)
    
    excel_synchronization_repo.bulk_create(db_session, [sale_in, sale_in2, sale_in3])
    
    res = get_pending_sales(db_session, target_date=date(2023, 10, 5))
    assert len(res) == 2
    
    res2 = get_pending_sales(db_session, target_date=date(2023, 10, 6))
    assert len(res2) == 1
    
    # Check default today (should be 0 since we only added for past dates)
    res3 = get_pending_sales(db_session)
    assert len(res3) == 0

def test_find_candidates_by_ean(db_session):
    blueprint1_id, batch_id = create_base_data(db_session)
    
    a1 = Article(article_blueprint_id=blueprint1_id, batch_id=batch_id, supplier_code="VEND-ABC", ean="1234567890123", colors=["Red"], status=ArticleStatus.AVAILABLE)
    db_session.add(a1)
    db_session.commit()
    db_session.refresh(a1)
    
    candidates = find_candidates(db_session, search_code="1234567890123", is_exchange=False)
    assert len(candidates) == 1
    assert candidates[0]["article_id"] == str(a1.id)

def test_ingest_captures_raw_code(db_session, monkeypatch):
    from src.core.config import settings
    monkeypatch.setattr(settings, "EXCEL_CODE_COLUMN", "N") # Zero based index 13
    
    rows_data = [
        # A      B     C     D     E     F     G     H     I     J     K     L     M     N
        [100.0, 90.0, None, None, None, None, None, None, None, None, None, None, None, "1234567890123"]
    ]
    
    file_path = create_mock_excel(rows_data, has_headers_at_699=False)
    target_date = date(2023, 11, 1)
    
    try:
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 1
        
        records = excel_synchronization_repo.get_pending_records(db_session, target_date=target_date)
        assert len(records) == 1
        assert records[0].raw_article_code == "1234567890123"
        assert records[0].status == ExcelSaleStatus.ORPHAN
    finally:
        os.remove(file_path)

def test_ingest_no_raw_code_column(db_session, monkeypatch):
    from src.core.config import settings
    monkeypatch.setattr(settings, "EXCEL_CODE_COLUMN", "")
    
    rows_data = [
        # A      B     C     D     E     F     G     H     I     J     K     L     M     N
        [100.0, 90.0, None, None, None, None, None, None, None, None, None, None, None, "1234567890123"]
    ]
    
    file_path = create_mock_excel(rows_data, has_headers_at_699=False)
    target_date = date(2023, 11, 2)
    
    try:
        inserted_ids = ingest_daily_sales(db_session, file_path, target_date)
        assert len(inserted_ids) == 1
        
        records = excel_synchronization_repo.get_pending_records(db_session, target_date=target_date)
        assert len(records) == 1
        assert records[0].raw_article_code is None
        assert records[0].status == ExcelSaleStatus.ORPHAN
    finally:
        os.remove(file_path)

def test_pending_sales_includes_raw_code(db_session):
    sale_in = ExcelSaleCreate(date=date(2023, 11, 3), excel_row_index=1, status=ExcelSaleStatus.ORPHAN, is_exchange=False, raw_article_code="VEND-XYZ")
    excel_synchronization_repo.create(db_session, obj_in=sale_in)
    
    from src.services.excel_synchronization_service import get_pending_sales
    res = get_pending_sales(db_session, target_date=date(2023, 11, 3))
    assert len(res) == 1
    assert res[0]["raw_article_code"] == "VEND-XYZ"

def test_get_daily_sales_returns_all_statuses(db_session):
    from src.services.excel_synchronization_service import get_daily_sales
    sale1 = ExcelSaleCreate(date=date(2023, 11, 4), excel_row_index=1, status=ExcelSaleStatus.ORPHAN, is_exchange=False)
    sale2 = ExcelSaleCreate(date=date(2023, 11, 4), excel_row_index=2, status=ExcelSaleStatus.RECONCILED, is_exchange=False)
    sale3 = ExcelSaleCreate(date=date(2023, 11, 4), excel_row_index=3, status=ExcelSaleStatus.UNPROCESSABLE, is_exchange=False)
    excel_synchronization_repo.bulk_create(db_session, [sale1, sale2, sale3])
    
    res = get_daily_sales(db_session, target_date=date(2023, 11, 4))
    assert len(res) == 3
    assert res[0]["status"] == ExcelSaleStatus.ORPHAN
    assert res[1]["status"] == ExcelSaleStatus.RECONCILED
    assert res[2]["status"] == ExcelSaleStatus.UNPROCESSABLE

from fastapi.testclient import TestClient
from src.main import app

def test_daily_sales_route_200(db_session):
    client = TestClient(app)
    response = client.get("/api/v1/excel-synchronization/daily-sales?target_date=2023-11-04")
    assert response.status_code == 200
    assert response.json()["status"] == "success"

def test_vendite_route_200(db_session):
    client = TestClient(app)
    response = client.get("/vendite")
    assert response.status_code == 200
    assert b"Vendite - La Triestina" in response.content

def test_get_daily_aggregations(db_session):
    from src.schemas.excel_sales import ExcelSaleCreate, ExcelSaleStatus
    from src.services.excel_synchronization_service import get_daily_aggregations
    from src.repositories.excel_synchronization_repo import excel_synchronization_repo
    from datetime import date, timedelta
    
    test_date = date(2024, 1, 4) # A Thursday
    hist_date_1 = test_date - timedelta(days=7) # Thursday
    hist_date_2 = test_date - timedelta(days=1) # Wednesday
    
    sales = [
        # Today
        ExcelSaleCreate(date=test_date, excel_row_index=1, excel_file_column=0, status=ExcelSaleStatus.ORPHAN, starting_price=240, selling_price=100, is_exchange=False, raw_article_code="A"),
        ExcelSaleCreate(date=test_date, excel_row_index=2, excel_file_column=0, status=ExcelSaleStatus.ORPHAN, starting_price=240, selling_price=150, is_exchange=True, raw_article_code="B"),
        ExcelSaleCreate(date=test_date, excel_row_index=3, excel_file_column=1, status=ExcelSaleStatus.ORPHAN, starting_price=210, selling_price=200, is_exchange=False, raw_article_code="C"),
        
        # Hist 1 (Same DOW)
        ExcelSaleCreate(date=hist_date_1, excel_row_index=4, excel_file_column=0, status=ExcelSaleStatus.ORPHAN, starting_price=240, selling_price=200, is_exchange=False, raw_article_code="A"),
        
        # Hist 2 (Diff DOW)
        ExcelSaleCreate(date=hist_date_2, excel_row_index=5, excel_file_column=0, status=ExcelSaleStatus.ORPHAN, starting_price=240, selling_price=220, is_exchange=False, raw_article_code="A"),
    ]
    
    excel_synchronization_repo.bulk_create(db_session, sales)
    db_session.commit()
    
    res = get_daily_aggregations(db_session, test_date)
    
    assert "columns" in res
    assert "pl_column" in res
    assert "pnl_summary" in res
    
    cols = {c["col_index"]: c for c in res["columns"]}
    
    col_0 = cols[0]
    assert col_0["total_revenue"] == 250
    assert col_0["total_list_price"] == 480
    assert col_0["discount"] == 230
    assert col_0["total_sales"] == 2
    assert col_0["avg_price"] == 125
    assert "target_date_dow" in res
    
    pnl = res["pnl_summary"]
    
    # Today's Net Discounted should be 250 + 200 = 450
    assert pnl["net_discounted"]["today"] == 450
    
    # DOW Avg Net Discounted:
    # 2 days (test_date, hist_date_1)
    # Total revenue for DOW = 450 (today) + 200 (hist1) = 650
    # Avg = 325
    assert pnl["net_discounted"]["dow_avg"] == 325
    
    # General Avg Net Discounted:
    # 3 days total
    # Total revenue = 450 + 200 + 220 = 870
    # Avg = 870 / 3 = 290
    assert pnl["net_discounted"]["general_avg"] == 290
    
    # Diff checks
    assert pnl["net_discounted"]["dow_diff"] == 125  # 450 - 325
    assert pnl["net_discounted"]["general_diff"] == 160  # 450 - 290
    
    # Constant check
    assert pnl["sg_media_giorno"]["is_constant"] is True
    assert pnl["sg_media_giorno"]["today"] == 500
    assert pnl["sg_media_giorno"]["dow_avg"] == 500
    assert pnl["sg_media_giorno"]["dow_diff"] == 0
